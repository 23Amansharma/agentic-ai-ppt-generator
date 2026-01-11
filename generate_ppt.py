from pptx import Presentation
from pptx.util import Pt, Inches

prs = Presentation()

def set_paragraph_format(paragraph, size=18, bold=False):
    font = paragraph.font
    font.size = Pt(size)
    font.bold = bold

def add_title_slide(title, subtitle):
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    subtitle_tf = slide.placeholders[1].text_frame
    subtitle_tf.clear()
    p = subtitle_tf.paragraphs[0]
    p.text = subtitle
    set_paragraph_format(p, size=14, bold=False)

def add_bullet_slide(title, bullet_points):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    content = slide.placeholders[1].text_frame
    content.clear()
    for i, point in enumerate(bullet_points):
        p = content.add_paragraph() if i>0 else content.paragraphs[0]
        p.text = point
        p.level = 0
        set_paragraph_format(p, size=18, bold=(i==0))

def add_section_slide(title):
    slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    title_shape.text = title
    set_paragraph_format(slide.shapes.title.text_frame.paragraphs[0], size=28, bold=True)

# --- Build improved deck ---
add_title_slide(
    "AGENTIC AI FOR INTELLIGENT TRAFFIC MANAGEMENT",
    "Satellite-Driven Congestion Control with Autonomous Agents\nTeam: Alok Gupta, Aman Sharma | Google Agentic AI Hackathon 2025"
)

add_section_slide("Agenda")
add_bullet_slide("Agenda", [
    "Problem Statement",
    "Our Agentic Solution",
    "System Architecture",
    "Pilot Results & Impact",
    "Next Steps"
])

add_bullet_slide("Problem — Urban Traffic Crisis", [
    "Reactive control causes long average response time (22 min)",
    "Satellite blind spots cause high downtime (68%) during monsoon",
    "Existing systems lack scalable, real-time adaptation"
])

add_bullet_slide("Agentic Solution Overview", [
    "Distributed autonomous agents coordinating via satellite telemetry",
    "Real-time local decision-making to reroute traffic and deploy resources",
    "Resilient to partial data loss using local sensing and agent negotiation"
])

add_bullet_slide("System Architecture", [
    "Satellite data ingestion → Agent mesh network → Local actuators",
    "Fallback: edge inference using onboard cameras and sensors",
    "Cloud orchestration for policy updates and analytics"
])

add_bullet_slide("Pilot Results & Impact", [
    "Simulated reduction in average response time from 22 → 6 minutes",
    "Estimated 35% reduction in congestion during peak hours",
    "Improved reliability under partial satellite outages"
])

add_bullet_slide("Next Steps & Ask", [
    "Pilot deployment in a busy urban corridor",
    "Partnerships with city traffic authorities and satellite providers",
    "Funding & compute resources for a 6-month field trial"
])

add_bullet_slide("Contact", [
    "Alok Gupta — alok@example.com",
    "Aman Sharma — aman@example.com",
    "Project repo and demo available upon request"
])

# Save updated deck
prs.save("Agentic_AI_Traffic_Management_Hackathon_2025.pptx")
